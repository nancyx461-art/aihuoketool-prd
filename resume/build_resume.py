#!/usr/bin/env python3
"""Generate the targeted resume (DOCX) for the Feishu healthcare KA sales role.

Usage:
    python3 resume/build_resume.py            # writes the .docx next to this script

Content lives in the data structures below so wording can be tuned without
touching the layout code.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent
OUT_NAME = "许增楠-医疗健康行业大客户销售-飞书.docx"

CJK_FONT = "微软雅黑"
LATIN_FONT = "Calibri"
ACCENT = RGBColor(0x2B, 0x4A, 0x6F)
TEXT = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
RULE = "C9D3DE"
BODY_PT = 9.5
SMALL_PT = 9

# --------------------------------------------------------------------------
# Content
# --------------------------------------------------------------------------

NAME = "许增楠"
TARGET = "求职意向：医疗健康行业-大客户销售"
CITY = "意向城市：深圳 / 广州"
CONTACT = [
    "15626104764",
    "15626104764@163.com",
    "企业管理硕士 · 市场营销方向",
    "英语 CET-6 / BEC 中级 · 粤语",
]

# Each strength: (lead-in, body). Keep the body to roughly two lines.
STRENGTHS = [
    (
        "医疗健康行业背景",
        "医药方向本科、医药营销方向硕士；强生医疗 1.5 年一线销售，罗氏制药、视源医疗显示实习；"
        "2024 年起以合伙人身份参与医疗器械渠道业务，开发 4 家医院、8 个产品入院。",
    ),
    (
        "医院与渠道客户开拓",
        "有从陌拜到准入、从科主任到副院长的医院客户开发经验，也有与渠道商合作谈判的经历，理解医院、厂家、渠道三方的诉求。",
    ),
    (
        "行业研究与表达",
        "主讲数字营销、营销渠道管理课程，主持医药内容营销课题并发表论文；曾负责卫星会与科室会宣讲，能独立完成行业分析并清晰讲解。",
    ),
    (
        "数字化与 AI 应用",
        "研究方向含数字营销与 AI 营销；日常使用 Claude、Cursor、Codex、WorkBuddy 等工具辅助研究与方案撰写，关注 AI 在企业服务中的落地。",
    ),
]

EXPERIENCE = [
    {
        "org": "广州南方学院 · 商学院",
        "title": "市场营销专任教师",
        "period": "2024.02 – 至今",
        "place": "广州",
        "bullets": [
            "主讲《数字营销》《营销渠道管理》《电子商务原理》课程，负责课程设计与学生指导。",
            "主持校级课题《中医药内容营销对消费意愿的影响机制研究》，成果发表于《商场现代化》（2026）。",
        ],
    },
    {
        "org": "医疗器械渠道业务（与前强生销售经理合伙）",
        "title": "合伙人（兼职）",
        "period": "2024 – 至今",
        "place": "广州",
        "bullets": [
            "负责新医院开发：对接副院长、科主任等决策人，完成 4 家医院准入、8 个产品入院。",
            "跨境医疗项目：带领团队与 4 家私立医院洽谈合作，建立 5 家合作渠道商。",
            "以渠道商身份参与经营，理解渠道商的经营逻辑与合作诉求。",
        ],
    },
    {
        "org": "强生（上海）医疗器材有限公司 · 能量外科与微创外科",
        "title": "销售培训生",
        "period": "2022.07 – 2024.01",
        "place": "广州",
        "bullets": [
            "区域生意规划：收集市场与竞品信息，制定区域销售策略并落地，超额完成业绩指标。",
            "重点客户开拓与维护：拜访核心客户，了解需求并提供针对性方案，建立长期客户关系。",
            "渠道管理：协助经销商制定行动计划，负责渠道人员产品培训，协同拜访重点客户，定期盘点渠道业绩。",
            "学术活动与跨部门协作：策划并落地 10+ 场学术活动（单场最高 40 万元）；配合市场部推进白皮书、手术指南编写；负责卫星会及科室会宣讲。",
        ],
    },
    {
        "org": "广州视源电子科技股份有限公司 · 医疗显示事业部",
        "title": "销售兼市场实习生",
        "period": "2021.01 – 2021.07",
        "place": "广州",
        "bullets": [
            "陌拜多家医院影像科开拓新市场，了解客户需求并传递产品价值；接待客户与渠道商参观并负责产品介绍。",
            "独立负责 4 部产品宣传视频与新闻稿的策划制作，对接媒体、设计等外部供应商。",
        ],
    },
    {
        "org": "上海罗氏制药有限公司 · 淋巴瘤产品组",
        "title": "区域市场部实习生",
        "period": "2019.03 – 2020.01",
        "place": "广州",
        "bullets": [
            "汇总区域销量、学术活动、专家反馈与竞品动态，输出销售月报与分析 PPT；协助区域市场活动的策划与执行。",
        ],
    },
]

PROJECT = {
    "org": "梅州泽山口腔医院 5 年战略规划",
    "title": "导师助理",
    "period": "2018.10 – 2019.10",
    "place": "",
    "bullets": [
        "对接医院决策层收集战略目标与竞争格局信息；研究口腔行业政策、技术与商业模式动态。",
        "对比竞品价目、测算人均业绩与人员饱和度，搭建盈亏平衡与客流分析表；完成 SWOT 分析与患者就诊流程调研，提出优化建议。",
    ],
}

RESEARCH = [
    "《中医药内容营销对消费意愿的影响机制：形而上与形而下内容的双路径作用》，《商场现代化》，2026（校级课题成果）。",
    "《透视直播电商解读营销的流变与永恒》，广东省社科联 2020 年岭南学术论坛会议论文，获评优秀论文。",
    "《国产动漫火爆，动漫 IP 如何打动消费者》，《国际品牌观察》，2020。",
]

EDUCATION = [
    {
        "org": "广东外语外贸大学",
        "title": "企业管理 · 硕士",
        "period": "2019.08 – 2022.06",
        "place": "广州",
        "bullets": [
            "研究方向：市场营销（品牌营销、医药营销）；连续三年获校级奖学金（一等 2 次、二等 1 次）。",
        ],
    },
    {
        "org": "广州中医药大学",
        "title": "国际经济与贸易（医药方向） · 本科",
        "period": "2015.09 – 2019.06",
        "place": "广州",
        "bullets": [
            "GPA 3.8，专业排名第 2；修读药理学、西医学概论等医药课程；连续四年获校级奖学金。",
        ],
    },
]

SKILLS = [
    ("AI 工具", "Claude、Cursor、OpenAI Codex、WorkBuddy，用于行业研究、方案撰写与数据整理。"),
    ("办公与分析", "Word / PowerPoint（熟练）、Excel（数据透视表、常用函数）、SPSS、AMOS、Xmind。"),
    ("语言", "英语（CET-6、BEC 中级）、粤语（良好）。"),
]

SUMMARY = (
    "医药专业出身，从药企市场实习、强生一线销售到高校研究与渠道实战，工作始终围绕医疗健康行业客户展开。"
    "希望加入飞书，把对医院、药企、渠道商的理解用于企业服务领域，做长期投入；入职后可全职投入，不再参与合伙业务经营。"
)

# --------------------------------------------------------------------------
# Layout helpers
# --------------------------------------------------------------------------


def set_run_font(run, size=BODY_PT, bold=False, color=TEXT):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)


def para(doc, before=0, after=2, line=1.2, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_bottom_border(paragraph, color=RULE, size=6, space=2):
    ppr = paragraph._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def usable_width(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def section_title(doc, text):
    p = para(doc, before=12, after=5, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=11.5, bold=True, color=ACCENT)
    add_bottom_border(p, color=RULE, size=6, space=3)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, lead=None, size=BODY_PT):
    p = para(doc, after=3, line=1.3)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.4)
    pf.first_line_indent = Cm(-0.4)
    pf.tab_stops.add_tab_stop(Cm(0.4))
    pf.keep_together = True
    r = p.add_run("•\t")
    set_run_font(r, size=size, color=ACCENT)
    if lead:
        r = p.add_run(lead + "　")
        set_run_font(r, size=size, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def entry_header(doc, org, title, period, place):
    p = para(doc, before=7, after=2, line=1.1)
    p.paragraph_format.tab_stops.add_tab_stop(usable_width(doc), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(org)
    set_run_font(r, size=10.5, bold=True)
    r = p.add_run("　" + title)
    set_run_font(r, size=BODY_PT, color=MUTED)
    right = period + ("　" + place if place else "")
    r = p.add_run("\t" + right)
    set_run_font(r, size=SMALL_PT, color=MUTED)
    return p


def entry(doc, item):
    entry_header(doc, item["org"], item["title"], item["period"], item["place"])
    paras = [bullet(doc, b) for b in item["bullets"]]
    # Short entries move to the next page as a block instead of leaving one bullet behind.
    if len(paras) <= 3:
        for p in paras[:-1]:
            p.paragraph_format.keep_with_next = True


def header_block(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    width = usable_width(doc)
    left, right = table.cell(0, 0), table.cell(0, 1)
    left.width = int(width * 0.58)
    right.width = int(width * 0.42)

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(NAME)
    set_run_font(r, size=24, bold=True, color=ACCENT)
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(TARGET)
    set_run_font(r, size=10.5, bold=True)
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(CITY)
    set_run_font(r, size=SMALL_PT + 0.5, color=MUTED)

    first = True
    for text in CONTACT:
        p = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, size=SMALL_PT, color=MUTED)

    sep = para(doc, before=4, after=0, line=1.0)
    add_bottom_border(sep, color="2B4A6F", size=10, space=1)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.9)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    header_block(doc)

    section_title(doc, "教育背景")
    for item in EDUCATION:
        entry(doc, item)

    section_title(doc, "核心优势")
    for lead, text in STRENGTHS:
        bullet(doc, text, lead=lead)

    section_title(doc, "工作经历")
    for item in EXPERIENCE:
        entry(doc, item)

    section_title(doc, "项目经历")
    entry(doc, PROJECT)

    section_title(doc, "论文与课题")
    for text in RESEARCH:
        bullet(doc, text, size=SMALL_PT)

    section_title(doc, "技能")
    for lead, text in SKILLS:
        bullet(doc, text, lead=lead)

    section_title(doc, "自我评价")
    p = para(doc, after=0, line=1.3)
    r = p.add_run(SUMMARY)
    set_run_font(r, size=BODY_PT)

    out = OUT_DIR / OUT_NAME
    doc.save(out)
    print(f"wrote {out}")
    return out


if __name__ == "__main__":
    build()
