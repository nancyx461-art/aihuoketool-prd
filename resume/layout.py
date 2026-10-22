"""Shared DOCX layout for the targeted resumes.

Each build script supplies a content dict and calls `render(content, out_path)`.
Content keys: NAME, TARGET, CITY, CONTACT, STRENGTHS, EXPERIENCE, PROJECTS,
RESEARCH, EDUCATION, SKILLS, SUMMARY, and optional SECTIONS (ordered list of
section keys to emit).
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK_FONT = "微软雅黑"
LATIN_FONT = "Calibri"
ACCENT = RGBColor(0x2B, 0x4A, 0x6F)
TEXT = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
RULE = "C9D3DE"
BODY_PT = 9.5
SMALL_PT = 9

DEFAULT_SECTIONS = [
    "EDUCATION", "STRENGTHS", "EXPERIENCE", "PROJECTS", "RESEARCH", "SKILLS", "SUMMARY",
]
SECTION_TITLES = {
    "EDUCATION": "教育背景",
    "STRENGTHS": "核心优势",
    "EXPERIENCE": "工作经历",
    "PROJECTS": "项目经历",
    "RESEARCH": "论文与课题",
    "SKILLS": "技能",
    "SUMMARY": "自我评价",
}


def set_run_font(run, size=BODY_PT, bold=False, color=TEXT):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)


def para(doc, before=0, after=2, line=1.2, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_bottom_border(paragraph, color=RULE, size=6, space=2):
    ppr = paragraph._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def usable_width(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def section_title(doc, text):
    p = para(doc, before=12, after=5, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=11.5, bold=True, color=ACCENT)
    add_bottom_border(p, color=RULE, size=6, space=3)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, lead=None, size=BODY_PT):
    p = para(doc, after=3, line=1.3)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.4)
    pf.first_line_indent = Cm(-0.4)
    pf.tab_stops.add_tab_stop(Cm(0.4))
    pf.keep_together = True
    r = p.add_run("•\t")
    set_run_font(r, size=size, color=ACCENT)
    if lead:
        r = p.add_run(lead + "　")
        set_run_font(r, size=size, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p


def entry_header(doc, org, title, period, place):
    p = para(doc, before=7, after=2, line=1.1)
    p.paragraph_format.tab_stops.add_tab_stop(usable_width(doc), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(org)
    set_run_font(r, size=10.5, bold=True)
    r = p.add_run("　" + title)
    set_run_font(r, size=BODY_PT, color=MUTED)
    right = period + ("　" + place if place else "")
    r = p.add_run("\t" + right)
    set_run_font(r, size=SMALL_PT, color=MUTED)
    return p


def entry(doc, item):
    entry_header(doc, item["org"], item["title"], item["period"], item["place"])
    paras = [bullet(doc, b) for b in item["bullets"]]
    # Short entries move to the next page as a block instead of leaving one bullet behind.
    if len(paras) <= 3:
        for p in paras[:-1]:
            p.paragraph_format.keep_with_next = True


def header_block(doc, c):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    width = usable_width(doc)
    left, right = table.cell(0, 0), table.cell(0, 1)
    left.width = int(width * 0.58)
    right.width = int(width * 0.42)

    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(c["NAME"])
    set_run_font(r, size=24, bold=True, color=ACCENT)
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(c["TARGET"])
    set_run_font(r, size=10.5, bold=True)
    p = left.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(c["CITY"])
    set_run_font(r, size=SMALL_PT + 0.5, color=MUTED)

    first = True
    for text in c["CONTACT"]:
        p = right.paragraphs[0] if first else right.add_paragraph()
        first = False
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, size=SMALL_PT, color=MUTED)

    sep = para(doc, before=4, after=0, line=1.0)
    add_bottom_border(sep, color="2B4A6F", size=10, space=1)


def render(c, out_path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.9)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    header_block(doc, c)

    for key in c.get("SECTIONS", DEFAULT_SECTIONS):
        data = c.get(key)
        if not data:
            continue
        section_title(doc, SECTION_TITLES[key])
        if key in ("EDUCATION", "EXPERIENCE", "PROJECTS"):
            for item in data:
                entry(doc, item)
        elif key in ("STRENGTHS", "SKILLS"):
            for lead, text in data:
                bullet(doc, text, lead=lead)
        elif key == "RESEARCH":
            for text in data:
                bullet(doc, text, size=SMALL_PT)
        elif key == "SUMMARY":
            p = para(doc, after=0, line=1.3)
            r = p.add_run(data)
            set_run_font(r, size=BODY_PT)

    doc.save(out_path)
    print(f"wrote {out_path}")
    return out_path
